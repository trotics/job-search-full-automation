"""Check a session's tracker writes against the rules.

Works with both trackers:
  - Spreadsheet: take snapshots with the `snapshot` command.
  - Artifact: list each collection with ArtifactData and save them together as one JSON file:
    {"companies": [...], "coverage": [...], "listings": [...], "answers": [...], "settings": [...]}

Commands:
  python tools/check_tracker.py snapshot my-files/job-search-tracker.xlsx backups/before.json
  python tools/check_tracker.py check backups/before.json backups/after.json --stage sweep
  python tools/check_tracker.py validate backups/after.json

Stages: sweep, fit, apply, outreach, prep, mailbox, resume, research, user.
Exit code 0 means every check passed.
"""
import json
import re
import sys

COLLECTIONS = ["companies", "coverage", "listings", "answers", "settings"]
STATUSES = ["To apply", "Applied", "Followed up", "Interview", "Offer", "Review fit", "On hold",
            "Rejected", "Closed", "expired", "filled"]
LOCKED = {"Applied", "Followed up", "Interview", "Offer"}
OUTREACH = ["managerName", "managerTitle", "howIdentified", "confidence", "profileUrl",
            "emailOrFormat", "connectionNote", "longMessage"]
NEW_LISTING_REQUIRED = ["id", "company", "title", "location", "url", "industry", "pay", "why",
                        "teaches", "priority", "status", "history", "postingText"]
SECRET_PATTERNS = [
    (re.compile(r"\b\d{3}-\d{2}-\d{4}\b"), "looks like a Social Security number"),
    (re.compile(r"password\s*[:=]", re.I), "looks like a stored password"),
    (re.compile(r"\b\d{13,19}\b"), "long number that could be a card or account number"),
]

# Which columns each stage may change on a listing that already existed.
STAGE_FIELDS = {
    "sweep": {"status", "history"},
    "fit": {"history"},
    "apply": {"status", "history", "postingText", "appliedDate"},
    "outreach": set(OUTREACH) | {"history"},
    "prep": {"history"},
    "mailbox": {"status", "history"},
    "research": set(),
    "resume": {"history"},
    "user": None,  # the user asked for a change directly: only the universal rules apply
}
# Which status moves each stage may make on an existing listing.
STAGE_MOVES = {
    "sweep": {("To apply", "expired"), ("To apply", "filled")},
    "fit": set(),
    "apply": {("To apply", "Applied"), ("To apply", "expired"), ("To apply", "filled")},
    "outreach": set(),
    "prep": set(),
    "mailbox": {(a, b) for a in ("Applied", "Followed up", "Interview")
                for b in ("Followed up", "Interview", "Rejected", "Offer") if a != b},
    "research": set(),
    "resume": set(),
    "user": None,
}


def flatten(row):
    """Accept {id, ...fields} or {id, data: {...}} (as some tools return)."""
    if isinstance(row, dict) and "data" in row and isinstance(row["data"], dict):
        out = dict(row["data"])
        out["id"] = row.get("id", out.get("id"))
        return out
    return dict(row)


def load(path):
    with open(path, encoding="utf-8") as f:
        raw = json.load(f)
    snap = {}
    for name in COLLECTIONS:
        rows = raw.get(name, [])
        if isinstance(rows, dict):  # {"docs": [...]} or {id: {...}}
            rows = rows.get("docs", [dict(v, id=k) for k, v in rows.items()])
        snap[name] = {str(r["id"]): r for r in (flatten(x) for x in rows)}
    return snap


def snapshot_xlsx(xlsx_path, out_path):
    from openpyxl import load_workbook
    wb = load_workbook(xlsx_path, data_only=True)
    out = {}
    for name in COLLECTIONS:
        rows = []
        if name in wb.sheetnames:
            ws = wb[name]
            header = [c.value for c in ws[1]]
            for r in ws.iter_rows(min_row=2, values_only=True):
                if all(v is None or str(v).strip() == "" for v in r):
                    continue
                rows.append({h: ("" if v is None else str(v)) for h, v in zip(header, r) if h})
        out[name] = rows
    with open(out_path, "w", encoding="utf-8") as f:
        json.dump(out, f, indent=1, ensure_ascii=False)
    print("Snapshot saved:", out_path, {k: len(v) for k, v in out.items()})


def norm(s):
    return re.sub(r"[^a-z0-9]", "", str(s or "").lower())


def validate(snap, problems, only_ids=None):
    names = {norm(c.get("company")): c.get("company") for c in snap["companies"].values()}
    for lid, l in snap["listings"].items():
        if only_ids is not None and lid not in only_ids:
            continue
        st = l.get("status", "")
        if st not in STATUSES:
            problems.append("listing %s: unknown status %r" % (lid, st))
        if l.get("company") and norm(l["company"]) in names and names[norm(l["company"])] != l["company"]:
            problems.append("listing %s: company spelled %r but target list says %r" % (lid, l["company"], names[norm(l["company"])]))
        if l.get("company") and norm(l["company"]) not in names:
            problems.append("listing %s: company %r is not on the target list" % (lid, l["company"]))
        for e in [x for x in str(l.get("history", "")).split("|") if x.strip()]:
            if not re.match(r"\s*\d{4}-\d{2}-\d{2}\s", e):
                problems.append("listing %s: history entry %r does not start with a YYYY-MM-DD date" % (lid, e.strip()[:40]))
        for col in ("history", "why", "postingText", "answer"):
            for pat, what in SECRET_PATTERNS:
                if pat.search(str(l.get(col, ""))):
                    problems.append("listing %s: %s %s" % (lid, col, what))
    seen = {}
    for lid, l in snap["listings"].items():
        key = (norm(l.get("company")), norm(l.get("reqId"))) if l.get("reqId") else None
        if key:
            if key in seen:
                problems.append("listings %s and %s: same company and req id" % (seen[key], lid))
            seen[key] = lid
    for aid, a in snap["answers"].items():
        for pat, what in SECRET_PATTERNS:
            if pat.search(str(a.get("answer", ""))) or pat.search(str(a.get("useWhen", ""))):
                problems.append("answer %s: %s" % (aid, what))
        if re.search(r"passw|\bpin\b|social security|routing", str(a.get("question", "")), re.I):
            problems.append("answer %s: the answer bank must never hold this kind of item" % aid)


def check(before, after, stage, problems):
    allowed_fields = STAGE_FIELDS[stage]
    allowed_moves = STAGE_MOVES[stage]
    b, a = before["listings"], after["listings"]

    for lid in b:
        if lid not in a:
            problems.append("listing %s was deleted" % lid)
    for lid, old in b.items():
        new = a.get(lid)
        if new is None:
            continue
        # Universal: yourNotes never written by a session.
        if str(old.get("yourNotes", "")) != str(new.get("yourNotes", "")):
            problems.append("listing %s: yourNotes changed (sessions never write it)" % lid)
        # Universal: history is append-only.
        oh, nh = str(old.get("history", "")), str(new.get("history", ""))
        if not nh.startswith(oh):
            problems.append("listing %s: history was rewritten, not appended to" % lid)
        changed = {k for k in set(old) | set(new)
                   if str(old.get(k, "")) != str(new.get(k, "")) and k not in ("yourNotes",)}
        if changed and nh == oh and stage != "user":
            problems.append("listing %s: changed %s with no history entry" % (lid, sorted(changed)))
        if nh.startswith(oh) and nh != oh:
            added = [e for e in nh[len(oh):].split("|") if e.strip()]
            if oh.strip() and not nh[len(oh):].lstrip().startswith("|"):
                problems.append("listing %s: new history entry not separated by ' | '" % lid)
            for e in added:
                if not re.match(r"\s*\d{4}-\d{2}-\d{2}\s", e):
                    problems.append("listing %s: history entry %r does not start with a YYYY-MM-DD date" % (lid, e.strip()[:40]))
        if allowed_fields is not None:
            extra = changed - allowed_fields
            if extra:
                problems.append("listing %s: stage %s may not change %s" % (lid, stage, sorted(extra)))
        os_, ns = old.get("status", ""), new.get("status", "")
        if os_ != ns and allowed_moves is not None and (os_, ns) not in allowed_moves:
            problems.append("listing %s: stage %s may not move status %r to %r" % (lid, stage, os_, ns))
        if stage in ("sweep", "fit") and os_ in LOCKED and changed:
            problems.append("listing %s: at %s, a %s may not touch it" % (lid, os_, stage))
        if stage == "apply" and ns == "Applied" and os_ != "Applied":
            if not str(new.get("postingText", "")).strip():
                problems.append("listing %s: set to Applied without postingText saved" % lid)
            if not str(new.get("appliedDate", "")).strip():
                problems.append("listing %s: set to Applied without appliedDate" % lid)
            if "upload:" not in str(new.get("history", ""))[len(str(old.get("history", ""))):].lower():
                problems.append("listing %s: set to Applied without the upload method in history" % lid)

    new_ids = set(a) - set(b)
    for lid in sorted(new_ids):
        l = a[lid]
        if stage not in ("sweep", "fit", "user"):
            problems.append("listing %s: stage %s may not add listings" % (lid, stage))
            continue
        missing = [c for c in NEW_LISTING_REQUIRED if not str(l.get(c, "")).strip()]
        if missing:
            problems.append("new listing %s: missing %s" % (lid, missing))
        if stage != "user" and l.get("status") != "To apply":
            problems.append("new listing %s: status must be To apply, not %r" % (lid, l.get("status")))
        if str(l.get("yourNotes", "")).strip():
            problems.append("new listing %s: yourNotes must be empty" % lid)
        if not re.fullmatch(r"[a-z0-9-]+", lid):
            problems.append("new listing %s: id must be lowercase letters, numbers and hyphens" % lid)

    # Other collections: only some stages write them.
    writers = {"companies": {"sweep", "research", "user"}, "coverage": {"sweep", "fit", "user"},
               "answers": {"apply", "user"}, "settings": {"user"}}
    for name, ok in writers.items():
        if before[name] != after[name] and stage not in ok:
            problems.append("%s changed, but stage %s does not write it" % (name, stage))
    for cid, old in before["coverage"].items():
        new = after["coverage"].get(cid)
        if new and str(old.get("detail", "")) and old.get("detail") != new.get("detail"):
            if str(old.get("detail", "")).strip() not in str(new.get("detail", "")):
                problems.append("coverage %s: earlier detail was not kept" % cid)

    validate(after, problems, only_ids=new_ids | {k for k in a if k in b and a[k] != b[k]})
    return new_ids


def main():
    args = sys.argv[1:]
    if not args:
        print(__doc__)
        sys.exit(2)
    cmd = args[0]
    if cmd == "snapshot":
        snapshot_xlsx(args[1], args[2])
        return
    problems = []
    if cmd == "validate":
        validate(load(args[1]), problems)
        label = "validate"
    elif cmd == "check":
        stage = args[args.index("--stage") + 1] if "--stage" in args else None
        if stage not in STAGE_FIELDS:
            sys.exit("Give --stage as one of: " + ", ".join(STAGE_FIELDS))
        before, after = load(args[1]), load(args[2])
        new_ids = check(before, after, stage, problems)
        changed = [k for k in after["listings"] if k in before["listings"] and after["listings"][k] != before["listings"][k]]
        print("Stage: %s | listings before %d, after %d | new %d | changed %d" % (
            stage, len(before["listings"]), len(after["listings"]), len(new_ids), len(changed)))
        label = "check"
    else:
        sys.exit("Unknown command " + cmd)
    for p in problems:
        print("  FAIL", p)
    print("RESULT: %s (%d problem%s)" % ("PASS" if not problems else "FAIL", len(problems), "" if len(problems) == 1 else "s"))
    sys.exit(1 if problems else 0)


if __name__ == "__main__":
    main()
