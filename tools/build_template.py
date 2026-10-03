"""Build the empty spreadsheet tracker: tracker/spreadsheet/job-search-tracker.xlsx plus one .csv per tab.

Usage:
  python tools/build_template.py                    (empty template)
  python tools/build_template.py data.json out.xlsx (fill from a JSON snapshot, for tests)
"""
import csv
import json
from datetime import datetime
import os
import sys

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from tracker_schema import CHOICES, COLUMNS, DEFAULT_SETTINGS  # noqa: E402

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT_DIR = os.path.join(ROOT, "tracker", "spreadsheet")

WIDE = {"detail": 60, "skipped": 50, "why": 50, "teaches": 40, "postingText": 60, "history": 70,
        "yourNotes": 40, "connectionNote": 50, "longMessage": 60, "answer": 50, "useWhen": 40,
        "question": 36, "industryOrder": 50, "excludedIndustries": 40, "roleFamilies": 50,
        "tierA": 40, "tierB": 40, "tierC": 40, "title": 36, "location": 36, "pay": 36, "company": 28,
        "careersSite": 40, "url": 40}

GUIDE = [
    ("Read me", ""),
    ("What this is", "Your job search tracker, for people who do not use the artifact tracker. Claude reads and writes it in sessions."),
    ("Before a session", "Close this file in Excel. Claude cannot safely write it while you have it open."),
    ("Your column", "yourNotes on the listings tab is yours. Claude never writes it."),
    ("History", "The history column on listings is a log. Claude only adds to the end. Do not delete entries."),
    ("Status", "Sweeps never change a listing at Applied, Followed up, Interview or Offer."),
    ("Never store", "Passwords, ID numbers, bank or card numbers. Not in any tab."),
    ("Column guide", "Every column is described in skill/job-search/references/tracker-columns.md."),
]


def build(data, path):
    wb = Workbook()
    # Fixed, neutral file properties: no build time, no computer user name.
    wb.properties.created = wb.properties.modified = datetime(2026, 1, 1)
    wb.properties.creator = "Job Search Full Automation"
    wb.properties.lastModifiedBy = None
    guide = wb.active
    guide.title = "read me"
    for i, (a, b) in enumerate(GUIDE, start=1):
        guide.cell(row=i, column=1, value=a).font = Font(bold=True)
        guide.cell(row=i, column=2, value=b).alignment = Alignment(wrap_text=True, vertical="top")
    guide.column_dimensions["A"].width = 20
    guide.column_dimensions["B"].width = 100

    head_fill = PatternFill("solid", fgColor="E9E4F3")
    user_fill = PatternFill("solid", fgColor="FFF4D6")
    for name, cols in COLUMNS.items():
        ws = wb.create_sheet(name)
        ws.append(cols)
        for j, col in enumerate(cols, start=1):
            c = ws.cell(row=1, column=j)
            c.font = Font(bold=True)
            c.fill = user_fill if col == "yourNotes" else head_fill
            ws.column_dimensions[get_column_letter(j)].width = WIDE.get(col, 16)
        ws.freeze_panes = "B2"
        for row in data.get(name, []):
            ws.append([row.get(c, "") for c in cols])
        for j, col in enumerate(cols, start=1):
            opts = CHOICES.get((name, col))
            if opts:
                dv = DataValidation(type="list", formula1='"%s"' % ",".join(opts), allow_blank=True)
                ws.add_data_validation(dv)
                L = get_column_letter(j)
                dv.add("%s2:%s2000" % (L, L))
        for r in ws.iter_rows(min_row=2):
            for c in r:
                c.alignment = Alignment(wrap_text=False, vertical="top")
    wb.save(path)


def write_csvs(data):
    for name, cols in COLUMNS.items():
        with open(os.path.join(OUT_DIR, name + ".csv"), "w", newline="", encoding="utf-8") as f:
            w = csv.writer(f)
            w.writerow(cols)
            for row in data.get(name, []):
                w.writerow([row.get(c, "") for c in cols])


def main():
    if len(sys.argv) == 3:
        with open(sys.argv[1], encoding="utf-8") as f:
            data = json.load(f)
        build(data, sys.argv[2])
        print("Wrote", sys.argv[2])
        return
    os.makedirs(OUT_DIR, exist_ok=True)
    data = {name: [] for name in COLUMNS}
    data["settings"] = [DEFAULT_SETTINGS]
    build(data, os.path.join(OUT_DIR, "job-search-tracker.xlsx"))
    write_csvs(data)
    print("Wrote tracker/spreadsheet/job-search-tracker.xlsx and %d .csv files" % len(COLUMNS))


if __name__ == "__main__":
    main()
