"""Safe writes to the spreadsheet tracker. Claude uses this in sessions; you can too.

Commands:
  python tools/sheet.py backup  tracker/my-tracker.xlsx
  python tools/sheet.py show    tracker/my-tracker.xlsx listings <id>
  python tools/sheet.py list    tracker/my-tracker.xlsx listings          (every row of a tab, as JSON)
  python tools/sheet.py apply   tracker/my-tracker.xlsx tracker/backups/<stamp>-changes.json

The changes file holds personal data: keep it in tracker/backups/, which git never shares.
It is a list of changes, applied in order:
  {"tab": "listings", "op": "add",    "id": "acme-r12", "fields": {...}}       new row (id must be new)
  {"tab": "listings", "op": "update", "id": "acme-r12", "fields": {"status": "expired"}}
  {"tab": "listings", "op": "history","id": "acme-r12", "entry": "2026-03-14 posting removed on the employer's site."}

Rules it enforces:
  - yourNotes is never written.
  - history is never replaced, only appended to with the "history" op, as " | <entry>".
  - every history entry starts with a YYYY-MM-DD date.
  - the file must be closed in Excel (an open file cannot be saved).
"""
import json
import os
import re
import shutil
import sys
from datetime import datetime

from openpyxl import load_workbook

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from tracker_schema import APPEND_ONLY, COLUMNS, USER_ONLY  # noqa: E402


def backup(path):
    folder = os.path.join(os.path.dirname(os.path.abspath(path)), "backups")
    os.makedirs(folder, exist_ok=True)
    out = os.path.join(folder, "tracker-backup-%s.xlsx" % datetime.now().strftime("%Y-%m-%d-%H%M"))
    shutil.copy2(path, out)
    print("Backup saved:", out)


def index(ws):
    header = [c.value for c in ws[1]]
    rows = {}
    for r in range(2, ws.max_row + 1):
        v = ws.cell(row=r, column=1).value
        if v not in (None, ""):
            rows[str(v)] = r
    return header, rows


def apply(path, changes):
    wb = load_workbook(path)
    for n, ch in enumerate(changes, start=1):
        tab, op, rid = ch["tab"], ch["op"], str(ch["id"])
        if tab not in COLUMNS:
            sys.exit("change %d: unknown tab %r" % (n, tab))
        ws = wb[tab]
        header, rows = index(ws)
        fields = ch.get("fields", {})
        for k in fields:
            if (tab, k) in USER_ONLY:
                sys.exit("change %d: %s.%s belongs to the user and is never written" % (n, tab, k))
            if k not in header:
                sys.exit("change %d: %s has no column %r" % (n, tab, k))
        if op == "add":
            if rid in rows:
                sys.exit("change %d: id %s already exists in %s" % (n, rid, tab))
            hist = fields.get("history", "")
            if (tab, "history") in APPEND_ONLY and hist and not re.match(r"\d{4}-\d{2}-\d{2}\s", hist):
                sys.exit("change %d: history must start with a YYYY-MM-DD date" % n)
            row = dict(fields, id=rid)
            ws.append([row.get(c, "") for c in header])
        elif op == "update":
            if rid not in rows:
                sys.exit("change %d: id %s not found in %s" % (n, rid, tab))
            for k, v in fields.items():
                if (tab, k) in APPEND_ONLY:
                    sys.exit("change %d: use the history op to add to %s" % (n, k))
                ws.cell(row=rows[rid], column=header.index(k) + 1, value=v)
        elif op == "history":
            if rid not in rows:
                sys.exit("change %d: id %s not found in %s" % (n, rid, tab))
            entry = ch["entry"].strip()
            if not re.match(r"\d{4}-\d{2}-\d{2}\s", entry):
                sys.exit("change %d: history entry must start with a YYYY-MM-DD date" % n)
            if "|" in entry:
                sys.exit("change %d: a history entry may not contain the | character (it separates entries); use / instead" % n)
            cell = ws.cell(row=rows[rid], column=header.index("history") + 1)
            old = str(cell.value or "")
            cell.value = (old + " | " + entry) if old.strip() else entry
        else:
            sys.exit("change %d: unknown op %r" % (n, op))
    try:
        wb.save(path)
    except PermissionError:
        sys.exit("Could not save. Close the tracker in Excel and try again. Nothing was written.")
    print("Applied %d change(s) to %s" % (len(changes), path))


def show(path, tab, rid):
    ws = load_workbook(path, data_only=True)[tab]
    header, rows = index(ws)
    if rid not in rows:
        sys.exit("not found")
    for j, h in enumerate(header, start=1):
        print("%s: %s" % (h, ws.cell(row=rows[rid], column=j).value or ""))


def list_tab(path, tab):
    ws = load_workbook(path, data_only=True)[tab]
    header = [c.value for c in ws[1]]
    rows = []
    for r in range(2, ws.max_row + 1):
        if ws.cell(row=r, column=1).value in (None, ""):
            continue
        rows.append({h: ws.cell(row=r, column=j).value or "" for j, h in enumerate(header, start=1) if h})
    print(json.dumps(rows, indent=1, ensure_ascii=False, default=str))


def main():
    a = sys.argv[1:]
    if len(a) < 2:
        print(__doc__)
        sys.exit(2)
    if a[0] == "backup":
        backup(a[1])
    elif a[0] == "show":
        show(a[1], a[2], a[3])
    elif a[0] == "list":
        list_tab(a[1], a[2])
    elif a[0] == "apply":
        with open(a[2], encoding="utf-8") as f:
            apply(a[1], json.load(f))
    else:
        sys.exit("unknown command")


if __name__ == "__main__":
    main()
